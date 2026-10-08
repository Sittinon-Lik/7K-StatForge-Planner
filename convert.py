import json
import os
from openpyxl import load_workbook


INPUT_FILE = "characters.xlsx"
OUTPUT_FILE = "characters.json"


def xlsx_to_json():
    # ตรวจสอบว่าไฟล์ Excel มีอยู่จริงไหม
    if not os.path.exists(INPUT_FILE):
        print(f"❌ ไม่พบไฟล์: {INPUT_FILE}")
        print(f"📁 Python กำลังหาไฟล์ใน:")
        print(os.getcwd())
        return

    try:
        # เปิด Excel
        workbook = load_workbook(INPUT_FILE, data_only=True)
        worksheet = workbook.active

        # อ่านหัวตาราง
        headers = [
            cell.value
            for cell in worksheet[1]
        ]

        data = []

        # อ่านข้อมูล
        for row in worksheet.iter_rows(
            min_row=2,
            values_only=True
        ):
            # ข้ามแถวว่าง
            if all(value is None for value in row):
                continue

            item = {}

            for header, value in zip(headers, row):
                if header is not None:
                    item[str(header)] = value

            data.append(item)

        # สร้าง JSON
        with open(
            OUTPUT_FILE,
            "w",
            encoding="utf-8"
        ) as file:
            json.dump(
                data,
                file,
                ensure_ascii=False,
                indent=2
            )

        # แสดงตำแหน่งไฟล์เต็มๆ
        output_path = os.path.abspath(OUTPUT_FILE)

        print()
        print("✅ แปลงไฟล์สำเร็จ!")
        print(f"📄 จำนวนข้อมูล: {len(data)}")
        print(f"📁 ไฟล์ JSON อยู่ที่:")
        print(output_path)

    except Exception as e:
        print()
        print("❌ เกิดข้อผิดพลาด:")
        print(e)


if __name__ == "__main__":
    xlsx_to_json()